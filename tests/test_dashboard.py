from datetime import date, datetime
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from urllib.error import HTTPError
from unittest.mock import patch
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook

from src.excel_builder import build_workbook
from src.fetcher import (
    _parse_mof_date,
    _remove_isolated_etf_price_spikes,
    add_market_notes,
    fetch_yahoo_market_data,
)
from src.main import (
    REPORT_END_MARKER,
    REPORT_START_MARKER,
    _render_readme_section,
    update_readme_dashboard,
)
from src.report_generator import (
    _build_market_context,
    _build_prompt,
    _validate_report_numbers,
    generate_daily_report,
)


class FetcherTests(unittest.TestCase):
    def test_parse_mof_era_dates(self):
        self.assertEqual(_parse_mof_date("S49.9.24").date(), date(1974, 9, 24))
        self.assertEqual(_parse_mof_date("R8.10.1").date(), date(2026, 10, 1))

    def test_yield_note_uses_basis_points(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 1), "Indicator": "JGB 10Y", "Symbol": "MOF:JGB:10Y", "Value": 1.0, "PriceType": "Yield"},
                {"Date": date(2026, 10, 2), "Indicator": "JGB 10Y", "Symbol": "MOF:JGB:10Y", "Value": 1.05, "PriceType": "Yield"},
            ]
        )
        noted = add_market_notes(observations)
        self.assertEqual(noted.iloc[1]["Note"], "前回観測から1日 / 変化: +5.0 bp")

    def test_equity_change_percentage_is_recalculated_from_observations(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 26), "Indicator": "日経平均株価", "Symbol": "^N225", "Value": 62189.65, "PriceType": "Price"},
                {"Date": date(2026, 10, 27), "Indicator": "日経平均株価", "Symbol": "^N225", "Value": 63178.32, "PriceType": "Price"},
            ]
        )
        noted = add_market_notes(observations)
        self.assertIn("(+1.59%)", noted.iloc[-1]["Note"])
        self.assertIn("前回観測から1日 / 変化:", noted.iloc[-1]["Note"])

    def test_isolated_reversing_etf_spike_is_excluded(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 3, 30), "Symbol": "1306.T", "Value": 278.70, "PriceType": "Price"},
                {"Date": date(2026, 3, 31), "Symbol": "1306.T", "Value": 37.14, "PriceType": "Price"},
                {"Date": date(2026, 4, 1), "Symbol": "1306.T", "Value": 280.0, "PriceType": "Price"},
            ]
        )
        cleaned = _remove_isolated_etf_price_spikes(observations)
        self.assertNotIn(37.14, cleaned["Value"].values)
        self.assertEqual(len(cleaned), 2)

    def test_yahoo_fetch_includes_topix_etf_separately(self):
        from unittest.mock import patch

        def fake_download(symbol, indicator):
            return pd.DataFrame(
                [{"Date": date(2026, 10, 1), "Indicator": indicator, "Symbol": symbol, "Value": 420.0, "PriceType": "Price"}]
            )

        with patch("src.fetcher._download_yahoo_symbol", side_effect=fake_download) as download:
            data = fetch_yahoo_market_data()

        self.assertIn("1306.T", data["Symbol"].values)
        self.assertIn(
            "TOPIX連動ETF (NEXT FUNDS TOPIX連動型上場投信)",
            data.loc[data["Symbol"] == "1306.T", "Indicator"].values,
        )
        download.assert_any_call("1306.T", "TOPIX連動ETF (NEXT FUNDS TOPIX連動型上場投信)")


class ExcelBuilderTests(unittest.TestCase):
    def test_workbook_contains_formula_dashboard_and_required_sheets(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 1), "Indicator": "日本国債2年利回り", "Symbol": "MOF:JGB:2Y", "Value": 1.0, "PriceType": "Yield", "Note": "初回取得"},
                {"Date": date(2026, 10, 2), "Indicator": "日本国債2年利回り", "Symbol": "MOF:JGB:2Y", "Value": 1.05, "PriceType": "Yield", "Note": "前日比: +5.0 bp"},
            ]
        )
        with TemporaryDirectory() as directory:
            output_path = build_workbook(observations, [], Path(directory) / "dashboard.xlsx")
            workbook = load_workbook(output_path, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Dashboard", "History", "Charts", "MarketData"])
            dashboard = workbook["Dashboard"]
            row = next(
                index
                for index in range(2, dashboard.max_row + 1)
                if dashboard.cell(index, 2).value == "MOF:JGB:2Y"
            )
            self.assertIn("MAXIFS", dashboard.cell(row, 3).value)
            self.assertIn("MarketData[Symbol]", dashboard.cell(row, 5).value)
            self.assertIn(f"(E{row}-F{row})*100", dashboard.cell(row, 8).value)


class ReadmeTests(unittest.TestCase):
    def test_unavailable_symbol_is_listed_explicitly(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 1), "Indicator": "日経平均株価", "Symbol": "^N225", "Value": 60000.0, "PriceType": "Price", "Note": "初回取得"},
            ]
        )
        section = _render_readme_section(observations, [])
        topix_row = next(line for line in section.splitlines() if "`^TPX`" in line)
        self.assertIn("| 種別 | 単位 | 注記 |", section)
        self.assertIn("| 価格 |", next(line for line in section.splitlines() if "`^N225`" in line))
        self.assertIn("データなし", topix_row)
        self.assertIn("価格", topix_row)
        etf_row = next(line for line in section.splitlines() if "`1306.T`" in line)
        self.assertIn("TOPIX連動ETF", etf_row)
        self.assertIn("NEXT FUNDS TOPIX連動型上場投信", etf_row)
        self.assertIn("| JPY/口 |", etf_row)

    def test_dashboard_marks_topix_etf_unit_and_index_substitution(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 27), "Indicator": "TOPIX連動ETF (NEXT FUNDS TOPIX連動型上場投信)", "Symbol": "1306.T", "Value": 420.0, "PriceType": "Price", "Note": "初回取得"},
            ]
        )
        section = _render_readme_section(observations, [])
        self.assertIn("JPY/口", section)
        self.assertIn("ETFは指数そのものではなく", section)

    def test_daily_report_replaces_its_own_readme_section(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 27), "Indicator": "日経平均株価", "Symbol": "^N225", "Value": 60000.0, "PriceType": "Price", "Note": "初回取得"},
            ]
        )
        with TemporaryDirectory() as directory:
            readme_path = Path(directory) / "README.md"
            readme_path.write_text(
                f"# test\n\nREADME本文を保持\n\n<!-- MARKET-DASHBOARD:START -->old dashboard<!-- MARKET-DASHBOARD:END -->\n\n"
                f"{REPORT_START_MARKER}\nold report\n{REPORT_END_MARKER}\n",
                encoding="utf-8",
            )
            fixed_jst = datetime(2026, 10, 27, 14, 36, tzinfo=ZoneInfo("Asia/Tokyo"))
            with patch("src.main.README_PATH", readme_path), patch(
                "src.main.datetime"
            ) as mocked_datetime:
                mocked_datetime.now.return_value = fixed_jst
                update_readme_dashboard(observations, [], "### 今日の市場をひとことで\nテスト")
            content = readme_path.read_text(encoding="utf-8")

        self.assertIn("## 市場ダッシュボード", content)
        self.assertIn("最終更新: 2026年10月27日 14:36 (JST)", content)
        self.assertIn("生成日時: 2026年10月27日 14:36 (JST)", content)
        self.assertIn("### 今日の市場をひとことで", content)
        self.assertNotIn("old report", content)
        self.assertIn("README本文を保持", content)

    def test_report_failure_replaces_stale_report_with_reason(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 27), "Indicator": "日経平均株価", "Symbol": "^N225", "Value": 60000.0, "PriceType": "Price", "Note": "初回取得"},
            ]
        )
        with TemporaryDirectory() as directory:
            readme_path = Path(directory) / "README.md"
            readme_path.write_text(
                f"# test\n{REPORT_START_MARKER}\n古いプレースホルダー\n{REPORT_END_MARKER}\n",
                encoding="utf-8",
            )
            with patch("src.main.README_PATH", readme_path):
                update_readme_dashboard(
                    observations, [], report_status="市場数値の検証に失敗: 37.14"
                )
            content = readme_path.read_text(encoding="utf-8")

        self.assertIn("レポートを掲載できませんでした。", content)
        self.assertIn("理由: 市場数値の検証に失敗: 37.14", content)
        self.assertNotIn("古いプレースホルダー", content)


class GeminiReportTests(unittest.TestCase):
    def test_context_labels_units_stale_yields_and_currency_returns(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 1), "Indicator": "MSCI ACWI", "Symbol": "ACWI", "Value": 100.0, "PriceType": "Price"},
                {"Date": date(2026, 10, 2), "Indicator": "MSCI ACWI", "Symbol": "ACWI", "Value": 110.0, "PriceType": "Price"},
                {"Date": date(2026, 10, 1), "Indicator": "ドル円 (USD/JPY)", "Symbol": "USDJPY=X", "Value": 150.0, "PriceType": "Price"},
                {"Date": date(2026, 10, 2), "Indicator": "ドル円 (USD/JPY)", "Symbol": "USDJPY=X", "Value": 160.0, "PriceType": "Price"},
                {"Date": date(2026, 9, 1), "Indicator": "日本国債10年利回り", "Symbol": "MOF:JGB:10Y", "Value": 3.0, "PriceType": "Yield"},
            ]
        )
        with patch("src.report_generator.fetch_related_news", return_value=[]):
            context = _build_market_context(observations)

        by_symbol = {series["symbol"]: series for series in context["markets"]}
        self.assertEqual(by_symbol["ACWI"]["unit"], "USD/口")
        self.assertEqual(by_symbol["USDJPY=X"]["unit"], "JPY/USD")
        self.assertEqual(by_symbol["MOF:JGB:10Y"]["lag_vs_latest_market_days"], 31)
        acwi_return = next(
            comparison
            for comparison in context["currency_return_comparisons"]
            if comparison["symbol"] == "ACWI"
        )
        self.assertEqual(acwi_return["usd_price_return_percent"], 10.0)
        self.assertEqual(acwi_return["indicative_jpy_return_percent"], 17.333)
        self.assertEqual(acwi_return["fx_effect_percentage_points"], 7.333)

    def test_numeric_validation_rejects_number_not_in_market_data(self):
        context = {
            "markets": [
                {
                    "latest": {"date": "2026-10-27", "value": 420.0},
                    "chart_period_summary": {"first_value": 278.7},
                }
            ]
        }
        with self.assertRaisesRegex(RuntimeError, "37.14"):
            _validate_report_numbers(
                "### 主な指標の動き\nETFの期間安値は37.14です。", context
            )

    def test_numeric_validation_checks_sections_after_news(self):
        context = {
            "markets": [{"latest": {"date": "2026-10-27", "value": 420.0}}]
        }
        report = (
            "### 国内外のニュースと市場の反応\n記事では37.14と報じられています。\n"
            "### オルカン・S&P 500を保有する人向けの見方\n"
            "期間リターンは37.14%です。"
        )
        with self.assertRaisesRegex(RuntimeError, "37.14"):
            _validate_report_numbers(report, context)

    def test_numeric_validation_accepts_data_and_derived_values(self):
        context = {
            "markets": [
                {
                    "latest": {"date": "2026-10-27", "value": 63178.32},
                    "change_since_previous_observation": {
                        "difference": 988.67,
                        "percent": 1.589,
                    },
                }
            ]
        }
        _validate_report_numbers(
            "### 主な指標の動き\n2026年10月27日、日経平均は63,178.32で、"
            "前回比+988.67（+1.59%）でした。S&P 500とも比較します。",
            context,
        )

    def test_report_uses_api_key_header_and_recent_data(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 26), "Indicator": "MSCI ACWI", "Symbol": "ACWI", "Value": 157.5, "PriceType": "Price"},
                {"Date": date(2026, 10, 27), "Indicator": "MSCI ACWI", "Symbol": "ACWI", "Value": 158.0, "PriceType": "Price"},
            ]
        )
        complete_report = "\n\n".join(
            [
                "### 今日の市場をひとことで\n市場は上昇しました。",
                "### 主な指標の動き\n株価は上昇しました。",
                "### 国内外のニュースと市場の反応\nニュースを確認しました。",
                "### オルカン・S&P 500を保有する人向けの見方\n長期では分散が重要です。",
                "### チャートで見る直近の流れ\n直近は上向きです。",
            ]
        )
        response_body = json.dumps(
            {"candidates": [{"content": {"parts": [{"text": complete_report}]}}]}
        ).encode("utf-8")

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def read(self):
                return response_body

        with patch("src.report_generator.urlopen", return_value=FakeResponse()) as mocked_urlopen:
            report = generate_daily_report(observations, api_key="test-api-key")

        request = mocked_urlopen.call_args.args[0]
        headers = {name.lower(): value for name, value in request.header_items()}
        payload = json.loads(request.data)
        prompt = payload["contents"][0]["parts"][0]["text"]
        self.assertEqual(report, complete_report)
        self.assertEqual(headers["x-goog-api-key"], "test-api-key")
        self.assertNotIn("test-api-key", request.full_url)
        self.assertIn('"value": 158.0', prompt)
        self.assertIn('"change_since_previous_observation"', prompt)
        self.assertIn('"difference": 0.5', prompt)
        self.assertIn('"percent": 0.317', prompt)
        self.assertEqual(payload["tools"], [{"google_search": {}}])
        self.assertEqual(payload["generationConfig"]["maxOutputTokens"], 8192)
        self.assertEqual(
            payload["generationConfig"]["thinkingConfig"], {"thinkingBudget": 1024}
        )
        self.assertIn("オルカン", prompt)
        self.assertIn("^TPX", prompt)
        self.assertIn("1306.T", prompt)

    def test_empty_gemini_response_is_reported_clearly(self):
        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def read(self):
                return b'{"candidates": []}'

        observations = pd.DataFrame(
            [{"Date": date(2026, 10, 27), "Indicator": "MSCI ACWI", "Symbol": "ACWI", "Value": 158.0, "PriceType": "Price"}]
        )
        with patch("src.report_generator.urlopen", return_value=FakeResponse()):
            with self.assertRaisesRegex(RuntimeError, "no report candidates"):
                generate_daily_report(observations, api_key="test-api-key")

    def test_missing_api_key_skips_report_generation(self):
        with patch.dict(os.environ, {"GEMINI_API_KEY": ""}), patch(
            "src.report_generator.urlopen"
        ) as mocked_urlopen:
            report = generate_daily_report(pd.DataFrame())
        self.assertIsNone(report)
        mocked_urlopen.assert_not_called()

    def test_prompt_requires_substantive_sections(self):
        prompt = _build_prompt({"markets": [{"indicator": "日経平均株価", "symbol": "^N225", "recent_observations": [{"date": "2026-10-27", "value": 63000.0}]}]})
        self.assertIn("各見出しの後に", prompt)
        self.assertIn("3〜5文", prompt)
        self.assertIn("1000〜1600文字", prompt)
        self.assertIn("国内と海外の材料", prompt)
        self.assertIn("為替", prompt)
        self.assertIn("一次データにない市場数値", prompt)
        self.assertIn("因果関係", prompt)
        self.assertIn("TOPIX連動ETF", prompt)
        self.assertIn("ニュースは検索で記事本文を確認できた事実のみ", prompt)
        self.assertIn("currency_return_comparisons", prompt)

    def test_prompt_mentions_optional_related_news(self):
        prompt = _build_prompt({
            "markets": [{"indicator": "日経平均株価", "symbol": "^N225", "recent_observations": [{"date": "2026-10-27", "value": 63000.0}]}],
            "related_news": [{"title": "日銀会合観測", "source": "Reuters"}],
        })
        self.assertIn("関連ニュース", prompt)
        self.assertIn("海外", prompt)
        self.assertIn("国内", prompt)

    def test_grounding_sources_are_included_in_report(self):
        observations = pd.DataFrame(
            [{"Date": date(2026, 10, 27), "Indicator": "MSCI ACWI", "Symbol": "ACWI", "Value": 158.0, "PriceType": "Price"}]
        )
        response_body = json.dumps(
            {
                "candidates": [
                    {
                        "content": {
                            "parts": [
                                {
                                    "text": "\n\n".join(
                                        [
                                            "### 今日の市場をひとことで\n市場は上昇しました。",
                                            "### 主な指標の動き\n株価は上昇しました。",
                                            "### 国内外のニュースと市場の反応\nニュースを確認しました。",
                                            "### オルカン・S&P 500を保有する人向けの見方\n長期では分散が重要です。",
                                            "### チャートで見る直近の流れ\n直近は上向きです。",
                                        ]
                                    )
                                }
                            ]
                        },
                        "groundingMetadata": {
                            "groundingChunks": [
                                {"web": {"title": "市場ニュース", "uri": "https://example.com/news"}},
                                {"web": {"title": "市場ニュース", "uri": "https://example.com/news"}},
                            ]
                        },
                    }
                ]
            }
        ).encode("utf-8")

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def read(self):
                return response_body

        with patch("src.report_generator.urlopen", return_value=FakeResponse()):
            report = generate_daily_report(observations, api_key="test-api-key")

        self.assertIn("### 参照したニュース・資料", report)
        self.assertEqual(report.count("https://example.com/news"), 1)

    def test_incomplete_report_is_rejected(self):
        observations = pd.DataFrame(
            [{"Date": date(2026, 10, 27), "Indicator": "MSCI ACWI", "Symbol": "ACWI", "Value": 158.0, "PriceType": "Price"}]
        )
        response_body = json.dumps(
            {"candidates": [{"content": {"parts": [{"text": "### 今日の市場をひとことで\n市場は上昇しました。"}]}}]}
        ).encode("utf-8")

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def read(self):
                return response_body

        with patch("src.report_generator.urlopen", return_value=FakeResponse()):
            with self.assertRaisesRegex(RuntimeError, "incomplete report"):
                generate_daily_report(observations, api_key="test-api-key")

    def test_generated_report_with_unverified_market_number_is_rejected(self):
        observations = pd.DataFrame(
            [
                {"Date": date(2026, 10, 26), "Indicator": "TOPIX連動ETF", "Symbol": "1306.T", "Value": 418.4, "PriceType": "Price"},
                {"Date": date(2026, 10, 27), "Indicator": "TOPIX連動ETF", "Symbol": "1306.T", "Value": 420.0, "PriceType": "Price"},
            ]
        )
        report_text = "\n\n".join(
            [
                "### 今日の市場をひとことで\nETFは上昇しました。",
                "### 主な指標の動き\n期間安値は37.14でした。",
                "### 国内外のニュースと市場の反応\nニュースは確認できませんでした。",
                "### オルカン・S&P 500を保有する人向けの見方\n長期では分散が重要です。",
                "### チャートで見る直近の流れ\n直近は上向きです。",
            ]
        )
        response_body = json.dumps(
            {"candidates": [{"content": {"parts": [{"text": report_text}]}}]}
        ).encode("utf-8")

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

            def read(self):
                return response_body

        with patch("src.report_generator.urlopen", return_value=FakeResponse()), patch(
            "src.report_generator.fetch_related_news", return_value=[]
        ):
            with self.assertRaisesRegex(RuntimeError, "37.14"):
                generate_daily_report(observations, api_key="test-api-key")

    def test_http_404_skips_report_generation_instead_of_crashing(self):
        observations = pd.DataFrame(
            [{"Date": date(2026, 10, 27), "Indicator": "MSCI ACWI", "Symbol": "ACWI", "Value": 158.0, "PriceType": "Price"}]
        )
        with patch(
            "src.report_generator.urlopen",
            side_effect=HTTPError("https://example.invalid", 404, "Not Found", hdrs=None, fp=None),
        ):
            report = generate_daily_report(observations, api_key="invalid-key")
        self.assertIsNone(report)


if __name__ == "__main__":
    unittest.main()