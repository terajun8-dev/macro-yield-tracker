"""Build a formula-driven Excel dashboard with normalized data and charts."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .config import INDICATORS, JGB_INDICATORS, WORKBOOK_PATH


HEADER_FILL = PatternFill("solid", fgColor="173F4F")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 24


def _add_table(sheet, name: str, reference: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _build_market_data_sheet(workbook: Workbook, data: pd.DataFrame) -> None:
    sheet = workbook.create_sheet("MarketData")
    headers = ["Indicator", "Symbol", "Date", "Value", "PriceType", "Note"]
    sheet.append(headers)
    for row in data[headers].itertuples(index=False, name=None):
        sheet.append(list(row))
    _style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet["C"][1:]:
        cell.number_format = "yyyy-mm-dd"
    for cell in sheet["D"][1:]:
        cell.number_format = "0.0000"
    widths = {"A": 25, "B": 18, "C": 13, "D": 14, "E": 12, "F": 52}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    if sheet.max_row > 1:
        _add_table(sheet, "MarketData", f"A1:F{sheet.max_row}")


def _build_dashboard_sheet(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "Dashboard"
    headers = [
        "Indicator", "Symbol", "Latest Date", "Previous Date", "Latest Value",
        "Previous Value", "PriceType", "Change", "Change %", "Note",
    ]
    sheet.append(headers)

    symbols = {**{symbol: name for name, symbol in INDICATORS.items()}}
    symbols.update({symbol: name for symbol, name in JGB_INDICATORS.items()})
    for symbol, indicator in symbols.items():
        sheet.append([indicator, symbol])
        row = sheet.max_row
        sheet.cell(row, 3, f'=IFERROR(MAXIFS(MarketData[Date],MarketData[Symbol],B{row}),"")')
        sheet.cell(
            row, 4,
            f'=IF(C{row}="","",IFERROR(MAXIFS(MarketData[Date],MarketData[Symbol],B{row},MarketData[Date],"<"&C{row}),""))',
        )
        sheet.cell(
            row, 5,
            f'=IF(C{row}="","",XLOOKUP(1,(MarketData[Symbol]=B{row})*(MarketData[Date]=C{row}),MarketData[Value],""))',
        )
        sheet.cell(
            row, 6,
            f'=IF(D{row}="","",XLOOKUP(1,(MarketData[Symbol]=B{row})*(MarketData[Date]=D{row}),MarketData[Value],""))',
        )
        sheet.cell(
            row, 7,
            f'=IFERROR(XLOOKUP(B{row},MarketData[Symbol],MarketData[PriceType],""),"")',
        )
        sheet.cell(
            row, 8,
            f'=IF(OR(E{row}="",F{row}=""),"",IF(G{row}="Yield",(E{row}-F{row})*100,E{row}-F{row}))',
        )
        sheet.cell(
            row, 9,
            f'=IF(OR(E{row}="",F{row}="",G{row}="Yield",F{row}=0),"",(E{row}-F{row})/F{row})',
        )
        sheet.cell(
            row, 10,
            f'=IF(E{row}="","データなし","前回取得日比: "&(C{row}-D{row})&"日 / "&IF(G{row}="Yield","前日比: "&TEXT(H{row},"+0.0;-0.0;0.0")&" bp","前日比: "&TEXT(I{row},"+0.00%;-0.00%;0.00%")))',
        )

    _style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 3).number_format = "yyyy-mm-dd"
        sheet.cell(row, 4).number_format = "yyyy-mm-dd"
        sheet.cell(row, 5).number_format = "0.0000"
        sheet.cell(row, 6).number_format = "0.0000"
        sheet.cell(row, 8).number_format = "+0.0000;-0.0000;0.0000"
        sheet.cell(row, 9).number_format = "+0.00%;-0.00%;0.00%"
    widths = [26, 18, 14, 14, 16, 16, 12, 14, 14, 52]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    _add_table(sheet, "DashboardTable", f"A1:J{sheet.max_row}")


def _build_history_sheet(workbook: Workbook, data: pd.DataFrame) -> None:
    sheet = workbook.create_sheet("History")
    indicators = list(dict.fromkeys(data["Indicator"].tolist()))
    sheet.append(["Date", *indicators])
    dates = sorted(data["Date"].drop_duplicates())
    values = {(row.Date, row.Indicator): row.Value for row in data.itertuples()}
    for market_date in dates:
        sheet.append([market_date, *[values.get((market_date, indicator)) for indicator in indicators]])
    _style_header(sheet)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet["A"][1:]:
        cell.number_format = "yyyy-mm-dd"
    sheet.column_dimensions["A"].width = 13
    for column in range(2, sheet.max_column + 1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 21


def _build_charts_sheet(workbook: Workbook, chart_paths: list[Path]) -> None:
    sheet = workbook.create_sheet("Charts")
    sheet.sheet_view.showGridLines = False
    next_row = 1
    for path in chart_paths:
        image = ExcelImage(str(path))
        scale = min(1, 900 / image.width)
        image.width = int(image.width * scale)
        image.height = int(image.height * scale)
        sheet.add_image(image, f"A{next_row}")
        next_row += (image.height + 19) // 20 + 2
    sheet.column_dimensions["A"].width = 120


def build_workbook(
    data: pd.DataFrame,
    chart_paths: list[Path],
    output_path: Path = WORKBOOK_PATH,
) -> Path:
    """Create Dashboard, History, Charts, and formula-source MarketData sheets."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    _build_dashboard_sheet(workbook)
    _build_history_sheet(workbook, data)
    _build_charts_sheet(workbook, chart_paths)
    _build_market_data_sheet(workbook, data)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output_path)
    return output_path